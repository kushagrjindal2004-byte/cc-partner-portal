import openpyxl
import json
import re

files = {
    '2026-07': r'C:/Users/HP/.gemini/antigravity/brain/32b9fde6-cb0e-4b70-813d-096952dbad80/.user_uploaded/media_1791369812526.xlsx',
    '2026-08': r'C:/Users/HP/.gemini/antigravity/brain/32b9fde6-cb0e-4b70-813d-096952dbad80/.user_uploaded/media_1791369812530.xlsx',
    '2026-09': r'C:/Users/HP/.gemini/antigravity/brain/32b9fde6-cb0e-4b70-813d-096952dbad80/.user_uploaded/media_1791369812489.xlsx',
}

all_banks = [
    {"id": "au", "name": "AU", "status": "FINAL", "display_order": 1},
    {"id": "axis", "name": "AXIS", "status": "FINAL", "display_order": 2},
    {"id": "axis_lic", "name": "AXIS LIC", "status": "FINAL", "display_order": 3},
    {"id": "hdfc", "name": "HDFC", "status": "FINAL", "display_order": 4},
    {"id": "idfc", "name": "IDFC", "status": "FINAL", "display_order": 5},
    {"id": "indus", "name": "INDUS", "status": "FINAL", "display_order": 6},
    {"id": "sbi", "name": "SBI", "status": "FINAL", "display_order": 7},
    {"id": "tata_neu", "name": "TATA NEU", "status": "FINAL", "display_order": 8},
    {"id": "kiwi", "name": "KIWI", "status": "FINAL", "display_order": 9},
    {"id": "bob", "name": "BOB", "status": "FINAL", "display_order": 10},
    {"id": "yes_zaggle", "name": "YES ZAGGLE", "status": "FINAL", "display_order": 11},
    {"id": "icici", "name": "ICICI BANK", "status": "FINAL", "display_order": 12},
    {"id": "rbl", "name": "RBL", "status": "FINAL", "display_order": 13}
]

known_managers = [
    'AZAM', 'INAYA', 'BHAVANI', 'BINOD MISHRA', 'DIVYAM', 'VINAY PANDEY', 'ALKESH SHUKLA', 'ZEESHAN HAIDER'
]

managers = {m: {'display_name': m.title(), 'pin_code': '1234'} for m in known_managers}
managers['AZAM']['display_name'] = 'AZAM'
managers['INAYA']['display_name'] = 'INAYA'
managers['BHAVANI']['display_name'] = 'BHAVANI'
managers['BINOD MISHRA']['display_name'] = 'BINOD MISHRA'
managers['DIVYAM']['display_name'] = 'DIVYAM'
managers['VINAY PANDEY']['display_name'] = 'VINAY PANDEY'
managers['ALKESH SHUKLA']['display_name'] = 'ALKESH SHUKLA'
managers['ZEESHAN HAIDER']['display_name'] = 'ZEESHAN HAIDER'

partner_to_manager = {}
issuances = {}

def get_bank_cols(ws):
    bank_cols = {}
    for c in range(4, ws.max_column+1, 2):
        val = str(ws.cell(1, c).value or '').upper()
        if not val or val == 'NONE':
            continue
        bank_id = None
        if 'AU ' in val or 'AU(' in val or val.startswith('AU ') or val == 'AU FINAL': bank_id = 'au'
        elif 'AXIS LIC' in val: bank_id = 'axis_lic'
        elif 'AXIS' in val: bank_id = 'axis'
        elif 'HDFC' in val: bank_id = 'hdfc'
        elif 'IDFC' in val: bank_id = 'idfc'
        elif 'INDUS' in val: bank_id = 'indus'
        elif 'SBI' in val: bank_id = 'sbi'
        elif 'TATA NEU' in val: bank_id = 'tata_neu'
        elif 'KIWI' in val: bank_id = 'kiwi'
        elif 'BOB' in val: bank_id = 'bob'
        elif 'YES' in val: bank_id = 'yes_zaggle'
        elif 'ICICI' in val: bank_id = 'icici'
        elif 'RBL' in val: bank_id = 'rbl'
        
        if bank_id:
            bank_cols[bank_id] = {'lm_col': c, 'cm_col': c+1}
    return bank_cols

for mcode, path in files.items():
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['MIS']
    bank_cols = get_bank_cols(ws)
    
    curr_mgr = None
    for r in range(4, ws.max_row+1):
        pname = ws.cell(r, 1).value
        if not pname: continue
        pname = str(pname).strip()
        if not pname or pname.upper() in ['PARTNERS', 'TOTAL', 'GRAND TOTAL']: continue
        
        is_mgr = False
        for km in known_managers:
            if pname.upper() == km or (pname.upper().startswith(km) and len(pname.split()) <= len(km.split())+1 and 'PVT' not in pname.upper() and 'SERVICES' not in pname.upper() and 'BOOSTER' not in pname.upper() and 'SOLUTIONS' not in pname.upper() and 'ENTERPRISES' not in pname.upper()):
                is_mgr = True
                curr_mgr = km
                break
        
        if is_mgr:
            continue
            
        if curr_mgr is None:
            curr_mgr = 'VINAY PANDEY'
            
        # Register partner
        partner_to_manager[pname] = curr_mgr
        
        # Read bank counts
        for bank_id, cols in bank_cols.items():
            lm_val = ws.cell(r, cols['lm_col']).value or 0
            cm_val = ws.cell(r, cols['cm_col']).value or 0
            try: lm_count = int(lm_val)
            except: lm_count = 0
            try: cm_count = int(cm_val)
            except: cm_count = 0
            
            issuances[(mcode, pname, bank_id)] = {
                'lm_count': lm_count,
                'cm_count': cm_count
            }

print(f'Total unique Channel Partners mapped: {len(partner_to_manager)}')
print(f'Total monthly issuance entries parsed: {len(issuances)}')

# Build October 2026 cycle with LM from September CM
for pname, mgr in partner_to_manager.items():
    for b in all_banks:
        bid = b['id']
        sept_entry = issuances.get(('2026-09', pname, bid), {'lm_count': 0, 'cm_count': 0})
        issuances[('2026-10', pname, bid)] = {
            'lm_count': sept_entry['cm_count'],
            'cm_count': 0
        }

# Generate SQL compatible with exact schema
sql_lines = []
sql_lines.append("-- ===========================================================================")
sql_lines.append("-- COMPLETE SUPABASE SEED SCRIPT FOR JULY, AUGUST, SEPTEMBER, AND OCTOBER 2026")
sql_lines.append("-- ===========================================================================\n")

# Safety Schema Adjustments
sql_lines.append("-- 1. Schema Safety Setup")
sql_lines.append("CREATE EXTENSION IF NOT EXISTS \"uuid-ossp\";")
sql_lines.append("ALTER TABLE managers ADD COLUMN IF NOT EXISTS pin_code TEXT DEFAULT '1234';")
sql_lines.append("ALTER TABLE banks ADD COLUMN IF NOT EXISTS display_order INT DEFAULT 0;\n")

# 1. Banks
sql_lines.append("-- 2. Master Banks")
for b in all_banks:
    sql_lines.append(f"INSERT INTO banks (id, name, status, display_order) VALUES ('{b['id']}', '{b['name']}', '{b['status']}', {b['display_order']}) ON CONFLICT (id) DO UPDATE SET name=EXCLUDED.name, status=EXCLUDED.status, display_order=EXCLUDED.display_order;")

# 2. Cycles
sql_lines.append("\n-- 3. Monthly Cycles")
cycles = [
    ('2026-07', 'July 2026', 'true', '2026-06'),
    ('2026-08', 'August 2026', 'true', '2026-07'),
    ('2026-09', 'September 2026', 'false', '2026-08'),
    ('2026-10', 'October 2026', 'false', '2026-09')
]
for c in cycles:
    sql_lines.append(f"INSERT INTO monthly_cycles (code, name, is_locked, prev_month_code) VALUES ('{c[0]}', '{c[1]}', {c[2]}, '{c[3]}') ON CONFLICT (code) DO UPDATE SET name=EXCLUDED.name, is_locked=EXCLUDED.is_locked, prev_month_code=EXCLUDED.prev_month_code;")

# 3. Managers, Partners & Issuances via PL/pgSQL block
sql_lines.append("\n-- 4. Managers, Channel Partners, and Card Issuances")
sql_lines.append("DO $$")
sql_lines.append("DECLARE")
sql_lines.append("    v_mgr UUID;")
sql_lines.append("    v_cp UUID;")
sql_lines.append("    v_ci UUID;")
sql_lines.append("BEGIN")

# Insert managers using IF check
for mkey in known_managers:
    m = managers[mkey]
    dname = m['display_name'].replace("'", "''")
    pin = m['pin_code']
    sql_lines.append(f"""
    SELECT id INTO v_mgr FROM managers WHERE name = '{dname}' LIMIT 1;
    IF v_mgr IS NULL THEN
        INSERT INTO managers (name, pin_code) VALUES ('{dname}', '{pin}') RETURNING id INTO v_mgr;
    ELSE
        UPDATE managers SET pin_code = '{pin}' WHERE id = v_mgr;
    END IF;
    """)

# Group partners by manager
mgr_grouped_cps = {}
for pname, mkey in partner_to_manager.items():
    if mkey not in mgr_grouped_cps:
        mgr_grouped_cps[mkey] = []
    mgr_grouped_cps[mkey].append(pname)

for mkey, cp_list in mgr_grouped_cps.items():
    dname = managers[mkey]['display_name'].replace("'", "''")
    sql_lines.append(f"\n    -- Manager: {dname}")
    sql_lines.append(f"    SELECT id INTO v_mgr FROM managers WHERE name = '{dname}' LIMIT 1;")
    
    for pname in cp_list:
        esc_pname = pname.replace("'", "''")
        sql_lines.append(f"""
    SELECT id INTO v_cp FROM channel_partners WHERE name = '{esc_pname}' AND manager_id = v_mgr LIMIT 1;
    IF v_cp IS NULL THEN
        SELECT id INTO v_cp FROM channel_partners WHERE name = '{esc_pname}' LIMIT 1;
        IF v_cp IS NULL THEN
            INSERT INTO channel_partners (name, manager_id) VALUES ('{esc_pname}', v_mgr) RETURNING id INTO v_cp;
        ELSE
            UPDATE channel_partners SET manager_id = v_mgr WHERE id = v_cp;
        END IF;
    END IF;
        """)
        
        # Monthly issuances for this partner across all 4 cycles
        for mcode in ['2026-07', '2026-08', '2026-09', '2026-10']:
            for b in all_banks:
                bid = b['id']
                entry = issuances.get((mcode, pname, bid))
                if entry and (entry['lm_count'] > 0 or entry['cm_count'] > 0):
                    lm = entry['lm_count']
                    cm = entry['cm_count']
                    sql_lines.append(f"""
    SELECT id INTO v_ci FROM card_issuances WHERE partner_id = v_cp AND bank_id = '{bid}' AND month_year = '{mcode}' LIMIT 1;
    IF v_ci IS NULL THEN
        INSERT INTO card_issuances (month_year, partner_id, bank_id, lm_count, cm_count) VALUES ('{mcode}', v_cp, '{bid}', {lm}, {cm});
    ELSE
        UPDATE card_issuances SET lm_count = {lm}, cm_count = {cm}, updated_at = now() WHERE id = v_ci;
    END IF;""")

sql_lines.append("\nEND $$;\n")

# Write to file
sql_content = "\n".join(sql_lines)
with open("update_all_3_months.sql", "w", encoding="utf-8") as f:
    f.write(sql_content)

print(f"Generated update_all_3_months.sql ({len(sql_content)} bytes)")
